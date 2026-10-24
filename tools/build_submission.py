"""Generate manifests, Postman collection, report, and a student-named ZIP.

Requires python-docx and Pillow for document/diagram generation only.
The microservices themselves have no third-party dependencies.
"""
from pathlib import Path
import json
import textwrap
import zipfile
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
STUDENT = '2025m203021'
NAME = 'Vattikunta Vishnu Priya'


def write(path, content):
    target = ROOT / path
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(content.strip() + '\n', encoding='utf-8')


def manifests():
    resources = []
    for service, image in [('inventory', 'inventory'), ('orders', 'orders')]:
        resources += [dict(apiVersion='v1', kind='PersistentVolumeClaim',
            metadata=dict(name=f'{service}-data', namespace='bookflow'),
            spec=dict(accessModes=['ReadWriteOnce'], resources=dict(requests=dict(storage='1Gi'))))]
        env = [dict(name='SERVICE_TOKEN', valueFrom=dict(secretKeyRef=dict(name='service-auth', key='token')))]
        if service == 'orders':
            env.append(dict(name='INVENTORY_URL', value='http://inventory:8000'))
        probe = lambda endpoint: dict(httpGet=dict(path=endpoint, port=8000), periodSeconds=5, timeoutSeconds=2, failureThreshold=6)
        resources += [dict(apiVersion='apps/v1', kind='Deployment',
            metadata=dict(name=service, namespace='bookflow'), spec=dict(
                replicas=1, strategy=dict(type='Recreate'), selector=dict(matchLabels=dict(app=service)),
                template=dict(metadata=dict(labels=dict(app=service)), spec=dict(
                    securityContext=dict(runAsNonRoot=True, runAsUser=10001, runAsGroup=10001, fsGroup=10001),
                    terminationGracePeriodSeconds=15,
                    containers=[dict(name=service, image=f'bookflow-{image}:1.0', imagePullPolicy='IfNotPresent',
                        ports=[dict(containerPort=8000)], env=env,
                        resources=dict(requests=dict(cpu='100m', memory='64Mi'), limits=dict(cpu='500m', memory='256Mi')),
                        startupProbe=probe('/health/ready'), readinessProbe=probe('/health/ready'), livenessProbe=probe('/health/live'),
                        securityContext=dict(allowPrivilegeEscalation=False, readOnlyRootFilesystem=True, capabilities=dict(drop=['ALL'])),
                        volumeMounts=[dict(name='data', mountPath='/data')])],
                    volumes=[dict(name='data', persistentVolumeClaim=dict(claimName=f'{service}-data'))]))))]
        resources += [dict(apiVersion='v1', kind='Service',
            metadata=dict(name=service, namespace='bookflow'),
            spec=dict(type='ClusterIP', selector=dict(app=service), ports=[dict(port=8000, targetPort=8000)]))]
    # JSON is valid YAML; separators let kubectl read the multi-document stream.
    write('k8s/workloads.yaml', '\n---\n'.join(json.dumps(r, indent=2) for r in resources))


def postman():
    items = []
    cases = [
        ('01 Inventory before order', 'GET', '{{inventory}}/books', None, [], 200, None),
        ('02 Place order', 'POST', '{{orders}}/orders', dict(customer=NAME, sku='BOOK-001', quantity=2),
            [{'key':'Idempotency-Key','value':'{{orderKey}}'}], 201, 'create'),
        ('03 Read order', 'GET', '{{orders}}/orders/{{orderId}}', None, [], 200, None),
        ('04 Replay same order', 'POST', '{{orders}}/orders', dict(customer=NAME, sku='BOOK-001', quantity=2),
            [{'key':'Idempotency-Key','value':'{{orderKey}}'}], 200, None),
        ('05 Inventory after order', 'GET', '{{inventory}}/books/BOOK-001', None, [], 200, None),
        ('06 Insufficient stock', 'POST', '{{orders}}/orders', dict(customer=NAME, sku='BOOK-001', quantity=1000),
            [{'key':'Idempotency-Key','value':'{{orderKey}}-reject'}], 201, 'reject'),
        ('07 Invalid quantity', 'POST', '{{orders}}/orders', dict(customer=NAME, sku='BOOK-001', quantity=0),
            [{'key':'Idempotency-Key','value':'{{orderKey}}-bad'}], 400, None),
        ('08 Conflicting key', 'POST', '{{orders}}/orders', dict(customer=NAME, sku='BOOK-001', quantity=3),
            [{'key':'Idempotency-Key','value':'{{orderKey}}'}], 409, None),
        ('09 Private endpoint without token', 'POST', '{{inventory}}/reservations',
            dict(order_id='unauthorized', sku='BOOK-001', quantity=1), [], 401, None),
    ]
    for name, method, url, body, headers, status, special in cases:
        req = dict(method=method, header=[dict(key='Content-Type',value='application/json')] + headers, url=url)
        if body is not None:
            req['body'] = dict(mode='raw', raw=json.dumps(body, indent=2), options=dict(raw=dict(language='json')))
        tests = [f'pm.test("HTTP {status}", () => pm.response.to.have.status({status}));']
        if special == 'create':
            tests += ['pm.collectionVariables.set("orderId", pm.response.json().id);',
                      'pm.test("Confirmed", () => pm.expect(pm.response.json().status).to.eql("CONFIRMED"));']
        if special == 'reject':
            tests += ['pm.test("Rejected", () => pm.expect(pm.response.json().status).to.eql("REJECTED"));']
        item = dict(name=name, request=req, event=[dict(listen='test',script=dict(type='text/javascript',exec=tests))])
        if special == 'create':
            item['event'].insert(0, dict(listen='prerequest',script=dict(type='text/javascript',exec=[
                'pm.collectionVariables.set("orderKey", "postman-" + pm.variables.replaceIn("{{$guid}}"));'])))
        items.append(item)
    collection = dict(info=dict(name=f'{STUDENT}_BookFlow', schema='https://schema.getpostman.com/json/collection/v2.1.0/collection.json'),
                      variable=[dict(key='inventory',value='http://127.0.0.1:8001'),dict(key='orders',value='http://127.0.0.1:8002'),
                                dict(key='orderKey',value=''),dict(key='orderId',value='')], item=items)
    write(f'postman/{STUDENT}_BookFlow.postman_collection.json', json.dumps(collection, indent=2))


README = r'''
# 2025m203021 — BookFlow

**Vattikunta Vishnu Priya · individual submission · SEZG583 Scalable Services**

Book ordering and stock reservation with two independently runnable Python microservices.
Prepared for the assignment due **12 November 2026**. The document filename supplied by the
course uses CCZG583; the report follows SEZG583, the course code inside that document.

## Start here

1. Read `2025m203021_BookFlow_Report.docx` or `2025m203021_BookFlow_Report.md`.
2. Install Python 3.13. Run `python tools/verify.py` to reproduce the tests and evidence.
3. In PowerShell, run `./tools/run-local.ps1`. Keep this terminal running.
4. In a second terminal in this folder, run `python tools/demo.py`.
5. Import `postman/2025m203021_BookFlow.postman_collection.json`; run requests in numbered order.
6. Complete `2025m203021_FINAL_CHECKLIST.md` before submitting the ZIP.

All commands in this package assume this folder is the working directory. If PowerShell
blocks a script, inspect it and run the commands manually or use an execution policy
permitted by your institution. Do not change a machine-wide policy for this assignment.

## Docker

Prerequisite: Docker Desktop running with Linux containers and Compose v2.

```powershell
$env:SERVICE_TOKEN = [guid]::NewGuid().ToString('N')
docker compose up --build -d
docker compose ps
python tools/demo.py
docker compose logs --tail 50
docker compose down
```

Named volumes preserve data after `down`. Stock is finite: repeated demonstrations can
exhaust it. Use an isolated test run for a fresh database. Remove named volumes only
when deliberately discarding demo data. Do not run local Python services at the same
time as containers using the same host ports.

## Minikube

Install Docker Desktop, kubectl and Minikube; enable your virtualization support.
Stop local Python services and Docker Compose before using ports 8001/8002 below.

```powershell
./tools/deploy-minikube.ps1
```

The script creates a separate `bookflow` Minikube profile and namespace, builds both
images inside Minikube, creates a random service Secret if none exists, applies PVCs,
ClusterIP Services and Deployments, and waits for both Deployments to become available.
It does not deploy to an unrelated current kubectl context.

Keep each forwarding command running in its own terminal:

```powershell
kubectl --context bookflow -n bookflow port-forward service/inventory 8001:8000
```

```powershell
kubectl --context bookflow -n bookflow port-forward service/orders 8002:8000
```

Run the demo or Postman against the same localhost URLs, then open the dashboard:

```powershell
minikube -p bookflow dashboard
kubectl --context bookflow -n bookflow get deployments,pods,services,pvc
kubectl --context bookflow -n bookflow describe deployment orders
kubectl --context bookflow -n bookflow logs deployment/orders --tail=50
kubectl --context bookflow -n bookflow get events --sort-by=.metadata.creationTimestamp
```

Select namespace **bookflow** in Dashboard. Capture both Deployments with 1/1 available,
both Pods Ready, both Services, bound PVCs, and deployment details. If a PVC is Pending,
check the default StorageClass and Minikube storage-provisioner/default-storageclass
addons. If images cannot be found, check the active profile and image tags. If startup
fails, inspect logs/events; both services must reference the same Secret token.

To demonstrate recovery, record an order and its stock level, then:

```powershell
kubectl --context bookflow -n bookflow rollout restart deployment/inventory
kubectl --context bookflow -n bookflow rollout status deployment/inventory --timeout=180s
```

Restart the Inventory port-forward after Pod replacement, query the same stock,
and replay the same Order POST with the same Idempotency-Key. Stock must not decrease
twice. `tools/demo.py` creates a NEW key each run, so use Postman request 04 to replay
an existing order. Stop the cluster after recording with `minikube -p bookflow stop`.
Changing images with the same tag requires rebuilding and restarting the Deployments.

## Independent service repositories

`inventory-service/` and `order-service/` are separate code roots. Each contains its own
runtime files, tests, Dockerfile, CI workflow and README, with no sibling imports.
When repository bundles are supplied, restore them with:

```powershell
git clone repository-bundles/2025m203021_inventory-service.bundle inventory-repo
git clone repository-bundles/2025m203021_order-service.bundle orders-repo
```

Create two empty GitHub repositories in your account. From each restored repository:

```powershell
git remote add origin https://github.com/YOUR_ACCOUNT/YOUR_REPOSITORY.git
git push -u origin main
```

Replace those values with the real destinations. Add both real URLs to the report.
GitHub publication has not been performed from this environment. Local CI definitions
run tests and build images on push; there is no automated deployment pipeline.

## Evidence and limitations

`evidence/api-results.json` contains actual calls to fresh local service processes.
`evidence/test-results.txt` contains actual test output. PNG response panels are
renderings of those captured responses, not Postman or Kubernetes screenshots.
Docker builds, Minikube, dashboard screenshots, GitHub uploads and a deployment
screen recording still require execution on a configured machine.

The teaching runtime uses ThreadingHTTPServer and a separate SQLite file per service.
The Minikube demo intentionally runs one replica per service with a Recreate strategy.
Do not increase replicas on a shared SQLite file or claim this configuration is
horizontally scalable. The report describes the production design with external
PostgreSQL databases, multiple API replicas and autoscaling. Local databases alone
do not establish high availability. There is no payment processing, customer login,
automatic pending-order reconciliation, cancellation, delivery or public ingress.

## Regenerate the document and ZIP

```powershell
python -m pip install -r tools/requirements-docs.txt
python tools/build_submission.py
```

Edit the Markdown report to add actual URLs, dashboard observations and video details
before regenerating. Rebuilding refreshes generated diagrams and API panels; it does
not manufacture missing deployment evidence. Your edits to an existing Markdown
report are retained. Source/tests run without the document-generation packages.
'''

REPORT = r'''
# BookFlow: Microservices Development and Deployment

## 1. Submission details

- Student ID: 2025m203021
- Student name: Vattikunta Vishnu Priya
- Group: Individual; no other members
- Course: SEZG583 Scalable Services
- Assignment: Microservices Development and Deployment
- Submission deadline: 12 November 2026
- Application name: BookFlow
- Submission filename: 2025m203021_BookFlow.zip

## 2. Application description and scope

BookFlow supports a small online bookstore. A customer browses books and places an
order for one book SKU and a positive integer quantity. Orders confirms the order
only after Inventory has durably reserved the stock. Rejected orders retain their
reason. Customers can query orders; staff can inspect available inventory.

The implemented vertical slice contains two services: Order and Inventory. A larger
deployment introduces Identity, Catalog, Payment, Fulfilment and Notification services.
Payment, cancellation, multi-item carts and shipping are outside this implementation.
Book titles and availability are co-located in Inventory only for the teaching slice;
the production Catalog owns searchable metadata. The demo customer field is a label,
not an authenticated identity. The API is intended for local assessment.

## 3. Business capabilities and service boundaries

| Business capability | Service | Owns | Implementation status |
|---|---|---|---|
| Order acceptance and order status | Order | Orders, client idempotency keys, state | Implemented |
| Stock availability and reservation | Inventory | Books, available stock, reservations | Implemented |
| Accounts and authorization | Identity | Users, credentials, roles | Proposed |
| Book search and metadata | Catalog | Titles, authors, search index | Proposed |
| Payment authorization and refunds | Payment | Payment attempts and provider references | Proposed |
| Packing and shipment tracking | Fulfilment | Shipments and delivery state | Proposed |
| Email/SMS status updates | Notification | Templates, delivery attempts | Proposed |

Boundaries follow business ownership. Order never queries Inventory's database;
Inventory never queries Order's database. Each has an independent schema, container,
repository and tests. An inventory schema change can be released independently while
preserving the REST contract. Both common.py copies are local to their repositories.

## 4. System operations: commands and queries

| Type | Operation | Owner | API / collaboration |
|---|---|---|---|
| Query | List available books | Inventory | GET /books |
| Query | Read one book's stock | Inventory | GET /books/{sku} |
| Command | Place order | Order | POST /orders; calls Inventory POST /reservations |
| Command | Reserve stock | Inventory | POST /reservations; service token required |
| Query | Read order | Order | GET /orders/{id} |
| Query | List up to 100 orders | Order | GET /orders |
| Query | Health / readiness | Both | GET /health/live and /health/ready |
| Command | Register / sign in | Identity | Proposed |
| Query | Search catalog | Catalog | Proposed |
| Command | Authorize / refund payment | Payment | Proposed, provider adapter |
| Command | Dispatch shipment | Fulfilment | Proposed, after confirmed payment |
| Command | Send status notification | Notification | Proposed, event consumer |

Business commands alter state; queries read it. Runtime health queries are operational,
separate from the bookstore use case. Stock replenishment is not implemented; seed
data is inserted only on first initialization and never overwrites existing stock.

## 5. Architecture and collaboration

![Implemented architecture](evidence/2025m203021_architecture.png)

The diagram shows the implemented topology. Solid arrows are synchronous HTTP/JSON;
each service owns its own persistent database. Kubernetes Services provide stable
DNS names inside the namespace. In local mode Order uses 127.0.0.1:8001; Docker and
Minikube use http://inventory:8000. Both services expose port 8000 inside containers;
host demo ports are 8001 and 8002.

### 5.1 Successful order sequence

1. Client POSTs customer, sku and quantity to Order with an Idempotency-Key header.
2. Order validates the fields, creates a PENDING order and commits locally.
3. Order POSTs the order ID, SKU and quantity to Inventory with X-Service-Token.
4. Inventory begins an immediate SQLite transaction. It replays an existing reservation
   or checks stock, inserts a reservation, and deducts stock atomically before commit.
5. Inventory returns RESERVED or REJECTED. Order stores CONFIRMED or REJECTED.
6. Client receives the durable order resource. A same-key same-payload retry returns
   the same order; a changed payload with the same key returns HTTP 409.

### 5.2 Failures, consistency and resilience

There is no cross-service ACID transaction. Inventory's reservation and stock update
are atomic locally. Order's initial PENDING write and later terminal update are local
transactions separated by HTTP. A two-second Inventory timeout bounds dependency wait.
Network failure or upstream 5xx/401 leaves Order PENDING and returns HTTP 503. The client
retries the same key. A crash after a successful reservation is safe because the retry
uses the persisted order ID and Inventory returns its original reservation without
deducting again. This is a reservation workflow with idempotent recovery, not a complete
payment/shipping saga. Exactly-once message delivery is not claimed.

Insufficient stock is a business result: Order creation returns HTTP 201 with status
REJECTED; Inventory's internal reservation returns 409. Unknown SKU also produces a
REJECTED order. A completed order replay returns 200. Quantity validation returns 400;
non-JSON content returns 415; unauthorized reservation requests return 401. A missing
resource returns 404. Overselling is prevented by BEGIN IMMEDIATE plus a nonnegative
stock constraint. Concurrent requests serialize their stock checks and updates.

PENDING recovery currently requires a client retry; an abandoned PENDING order can
remain unresolved. A future reconciliation worker would retry pending rows with backoff.
There is no reservation release or expiration: adding payment/cancellation requires
idempotent compensating release and reservation lifecycle rules.

### 5.3 Scalable production design

![Proposed scalable architecture](evidence/2025m203021_scalable_design.png)

Place a TLS gateway/load balancer before multiple stateless Order and Inventory API
replicas. Each service uses its own external PostgreSQL database with pooling,
backups, failover and unique idempotency constraints. Stock reservation must use row
locking or an atomic conditional UPDATE, preserving the same invariants across replicas.
Identity authenticates customers and service accounts; network policies and TLS protect
internal traffic. Catalog can scale search reads independently with an index/cache.

Introduce a broker for OrderConfirmed, PaymentAuthorized, ShipmentDispatched and
NotificationRequested events. An outbox in the publishing service commits state plus
events together; consumers use inbox/deduplication keys. Payment failure triggers an
idempotent release command through an orchestrated saga. Reads use pagination and
appropriate indexes. Apply rate limits, per-call timeouts, bounded retries with jitter,
circuit breakers, structured correlation IDs, metrics, tracing and alerting.

Horizontal Pod Autoscalers can target CPU/utilization once services use external
databases and a metrics provider. API replica counts can then vary independently;
database writes and hot SKUs remain bottlenecks requiring measurement and tuning.
Load-test latency, throughput, lock contention and error rate before choosing targets.
These production components are designed, not implemented or benchmarked in this submission.

The local SQLite demo uses one replica per service and a Recreate strategy to avoid
overlapping writers and ReadWriteOnce volume attachment conflicts. It demonstrates
service boundaries and container orchestration, not high availability or horizontal
scaling. Recreate updates have downtime. A future PostgreSQL migration enables rolling
updates and multiple replicas; increasing demo replicas is not a valid substitute.

## 6. Technology choices

| Concern | Choice | Reason / constraint |
|---|---|---|
| Language | Python 3.13 | Small independently testable implementation |
| REST runtime | ThreadingHTTPServer | Standard library; teaching/local assessment only |
| Data persistence | SQLite, separate file per service | Reproducible demo with atomic transactions |
| Communication | Synchronous REST/JSON, urllib timeout | Immediate reservation outcome; no broker dependency |
| Containerization | Docker, python:3.13-slim | Independent nonroot service images |
| Orchestration | Kubernetes on Minikube | Deployment, Service, Secret, probes and PVC demonstration |
| API exploration | Postman collection | Reproducible requests and response assertions |
| Tests | unittest and real HTTP subprocess tests | No external service dependency for unit suites |
| CI | GitHub Actions | Independent unit suites and image build per repository |

Python's documentation cautions that http.server is not recommended for production.
This assessment runtime should be replaced with an ASGI application/server and hardened
gateway in production. Images use a version tag rather than an immutable digest; a real
release pipeline should pin and scan approved base images. The optional CI configuration
is supplied; a remote Actions run and continuous deployment have not been performed.

## 7. Project structure and database ownership

```text
2025m203021_BookFlow/
  inventory-service/     independent service code repository
    app.py              routes, reservation logic and Inventory schema
    common.py           local HTTP validation and database helpers
    tests/test_inventory.py
    Dockerfile, requirements.txt, README.md
    .github/workflows/ci.yml
  order-service/         independent service code repository
    app.py              routes, order workflow and Order schema
    common.py           local HTTP validation and database helpers
    tests/test_orders.py
    Dockerfile, requirements.txt, README.md
    .github/workflows/ci.yml
  repository-bundles/    portable independent Git histories, when generated
  k8s/                  namespace and workload manifests
  postman/              importable collection with assertions
  tools/                local runner, demo, tests, Minikube script, packaging
  evidence/             actual local responses/test output, diagrams and API panels
  compose.yaml          separate service containers and volumes
  README.md             complete run and troubleshooting instructions
  2025m203021_BookFlow_Report.docx
  2025m203021_BookFlow_Report.md
  2025m203021_Demonstration_Script.md
  2025m203021_FINAL_CHECKLIST.md
```

Inventory schema: books(sku PK, title, available CHECK >= 0); reservations(order_id PK,
sku FK, quantity CHECK > 0, status). Order schema: orders(id PK, idempotency_key UNIQUE,
customer, sku, quantity CHECK > 0, status, reason, created_at). Timestamps are SQLite
CURRENT_TIMESTAMP in UTC. Customer labels and keys are limited to 120 characters;
quantities are integers from 1 to 1000. Booleans are not accepted as integers.
No database files are committed. Demo rows contain no real payment credentials.

## 8. Implementation steps

1. Model the order and reservation boundaries; identify each service's commands/queries.
2. Define JSON contracts and terminal/PENDING state transitions.
3. Implement service-owned schemas and transactional reservation logic.
4. Persist Order before HTTP and add same-key recovery for uncertain outcomes.
5. Add validation, service-token checks, health endpoints and upstream timeout handling.
6. Implement isolated unit suites and two-process HTTP integration checks.
7. Supply independent Dockerfiles, local Compose orchestration and Kubernetes manifests.
8. Prepare Postman requests, deployment instructions, architecture figures and evidence.

## 9. Execution instructions and deployment plan

See README.md for copyable PowerShell commands. For native execution run
./tools/run-local.ps1 and, from another terminal, python tools/demo.py. The local runner
uses persistent service-local data directories and stops its child processes on Ctrl+C.
For isolated verification use python tools/verify.py; it starts fresh temporary databases,
chooses free ports and stops all child processes after the run.

For Docker set SERVICE_TOKEN in the shell and run docker compose up --build -d.
Inventory must pass its readiness healthcheck before Compose starts Order. Named
volumes persist independent database files. Ports bind only to localhost.

For Kubernetes install the three prerequisites and run ./tools/deploy-minikube.ps1.
The script starts the bookflow profile, builds images inside that cluster, creates
the namespace and token Secret, applies the manifests and waits for deployments.
Each deployment requests 100m CPU/64Mi memory and limits 500m CPU/256Mi memory; these
are initial demo settings, not measured production capacity. Each receives a dedicated
1Gi ReadWriteOnce PVC. Containers run as UID 10001 with a read-only root filesystem,
privilege escalation disabled and Linux capabilities dropped; /data is writable.

Startup/readiness probes test the local database; liveness tests the HTTP process.
Order readiness deliberately does not depend on Inventory, so dependency outages
produce explicit API 503s without removing otherwise healthy Order pods from service.
The database readiness test verifies connection/query execution, not a full write test.
ClusterIP services stay internal. For assessment use separate localhost port-forwards;
no internet ingress is configured. Kubernetes Secrets are not automatically encrypted
at rest by this configuration. Do not publish live tokens.

## 10. Testing and captured evidence

The included evidence/test-results.txt records 14 passing unit tests, split evenly
between services, and a passing real two-service HTTP integration run. Coverage includes
idempotent replay, mismatched keys, insufficient stock, unknown SKU, invalid quantities,
timeout recovery, restart persistence and concurrent stock reservation.

In the HTTP run, BOOK-001 began with stock 10. Ordering quantity 2 produced CONFIRMED;
replaying the same key kept stock at 8. Twenty concurrent one-unit orders then produced
8 confirmations and 12 rejections, leaving stock at zero. Inventory was stopped, a
new order returned 503 and persisted PENDING, and restarting Inventory followed by
same-key retry confirmed that order. No performance or production-scale claim follows
from these correctness checks.

![Captured local API responses](evidence/2025m203021_api_evidence.png)

The panel above is rendered from actual evidence/api-results.json, with ephemeral local
ports. It is not a Postman screenshot or a Kubernetes deployment screenshot. Requests,
status codes and responses are available in that JSON for inspection. Import the
included collection to capture actual Postman screenshots of the same workflow.

### 10.1 Deployment evidence still to capture

Docker, kubectl, Minikube, GitHub CLI and a screen recording tool were not available
on the preparation machine's PATH. Container image builds and cluster execution were
not performed. The manifests and commands are provided for execution on a configured
machine. Replace this paragraph with actual run details after executing them; retain
honest distinctions between measured results and design proposals.

- Add a real Postman screenshot of order creation showing HTTP 201 and CONFIRMED.
- Add a same-key replay and stock query showing exactly one stock deduction.
- Add insufficient-stock and validation responses.
- Add docker compose ps showing two healthy service containers.
- Add Dashboard screenshots for the bookflow namespace: deployments, Pods and PVCs.
- Record observed replica readiness, restart counts, PVC binding, resource usage if
  metrics are enabled, and any warning Events. Do not invent successful observations.
- Add restart/recovery evidence and the demonstration video filename or approved link.

## 11. Kubernetes dashboard analysis guide

Inspect desired versus available replicas (target 1/1 each), readiness (1/1 containers),
image tags and restart counts. Check the two ClusterIP services select the correct
app labels and target port 8000. Confirm each PVC is Bound and service data survives a
Pod restart. Check Events for FailedScheduling, image pull errors, permission failures
or volume attachment errors; correlate failures with Pod logs and probe status.
If metrics are unavailable, report that fact rather than invent CPU/memory observations.
Explain why one local SQLite writer and Recreate updates are used. These are expected
inspection criteria; actual Dashboard observations are pending cluster execution.

## 12. Repository links and CI status

- Inventory GitHub repository: PENDING — upload the separate Inventory repository and add its actual URL.
- Order GitHub repository: PENDING — upload the separate Order repository and add its actual URL.

Both service roots are self-contained. Portable repository bundles, when included,
restore their independent commit histories without relying on the surrounding workspace.
Use the README's clone/push instructions to publish two distinct repositories.
Each CI workflow runs its own unit suite and builds its own Docker image on push/PR.
Successful local tests do not imply a successful remote Actions run. Optional automated
CD has not been implemented.

## 13. Demonstration recording

Follow 2025m203021_Demonstration_Script.md for a roughly 5–7 minute screen recording.
Show identity, architecture, independent repositories, container deployment, Minikube,
Dashboard, successful and rejected API calls, replay, and restart persistence.
Video status: PENDING — add 2025m203021_BookFlow_Demo.mp4 after recording a real deployment.
A script or rendered API panel does not fulfil the screen-recording requirement.

## 14. References

- Python HTTP server documentation: https://docs.python.org/3.13/library/http.server.html
- Python SQLite documentation: https://docs.python.org/3.13/library/sqlite3.html
- Docker Compose startup order: https://docs.docker.com/compose/how-tos/startup-order/
- Kubernetes probes: https://kubernetes.io/docs/tasks/configure-pod-container/configure-liveness-readiness-startup-probes/
- Kubernetes persistent volumes: https://kubernetes.io/docs/concepts/storage/persistent-volumes/
- Kubernetes autoscaling: https://kubernetes.io/docs/concepts/workloads/autoscaling/horizontal-pod-autoscale/
- Minikube image building/loading: https://minikube.sigs.k8s.io/docs/handbook/pushing/
- Minikube Dashboard: https://minikube.sigs.k8s.io/docs/handbook/dashboard/

The implementation was prepared specifically for this assignment with AI assistance.
References informed API/runtime and deployment choices. The student should review,
understand and adapt the implementation and disclose assistance as required by course
policy. No claim is made that externally hosted repositories, deployment screenshots
or a demonstration recording already exist.
'''


def diagrams():
    folder = ROOT / 'evidence'
    folder.mkdir(exist_ok=True)
    fontpath = Path('C:/Windows/Fonts/arial.ttf')
    def font(size):
        return ImageFont.truetype(str(fontpath), size) if fontpath.exists() else ImageFont.load_default(size=size)
    for production in [False, True]:
        canvas = Image.new('RGB', (1500, 910), '#f4f7fb')
        draw = ImageDraw.Draw(canvas)
        draw.text((45,30), 'BookFlow | ' + ('proposed scalable architecture' if production else 'implemented service architecture'), font=font(34), fill='#12243e')
        draw.text((45,85), '2025m203021  |  Vattikunta Vishnu Priya', font=font(22), fill='#4a607c')
        def box(rect, title, lines, color='#dceaff'):
            draw.rounded_rectangle(rect, radius=16, fill=color, outline='#7892b0', width=2)
            x,y,_,_ = rect
            draw.text((x+20,y+16), title, font=font(26), fill='#12243e')
            for i,line in enumerate(lines):
                draw.text((x+20,y+57+i*31), line, font=font(21), fill='#263b54')
        def arrow(start, end, label=None):
            draw.line((start, end), fill='#386393', width=4)
            x,y=end
            if start[1] == end[1]:
                draw.polygon([(x,y),(x-14,y-8),(x-14,y+8)],fill='#386393')
            else:
                draw.polygon([(x,y),(x-8,y-14),(x+8,y-14)],fill='#386393')
            if label:
                draw.text((min(start[0],end[0])+5,min(start[1],end[1])+12), label, font=font(18),fill='#386393')
        box((50,165,400,300),'Client / Postman',['Browse books','Place and query orders'])
        if production:
            box((560,165,960,300),'TLS gateway',['Identity / rate limiting','Routes to API replicas'])
            arrow((400,225),(560,225),'HTTPS')
            box((145,390,625,555),'Order API replicas',['Multiple independent Pods','PENDING -> CONFIRMED / REJECTED'])
            box((860,390,1340,555),'Inventory API replicas',['Atomic per-SKU reservation','Idempotent reservation ID'])
            arrow((760,300),(760,350))
            draw.line([(380,350),(1100,350)],fill='#386393',width=4)
            arrow((380,350),(380,390))
            arrow((1100,350),(1100,390))
            arrow((625,465),(860,465),'REST/JSON')
            box((145,625,625,755),'Order PostgreSQL',['Own schema, unique client key','Pool, backup, failover'],'#d9f2e5')
            box((860,625,1340,755),'Inventory PostgreSQL',['Own schema, stock row locks','Pool, backup, failover'],'#d9f2e5')
            arrow((380,555),(380,625))
            arrow((1100,555),(1100,625))
            draw.text((75,810),'Future: outbox -> broker -> Payment / Fulfilment / Notification; metrics + HPA',font=font(24),fill='#344d6b')
            draw.text((75,852),'Design proposal; these production components are not implemented in the local demo.',font=font(20),fill='#344d6b')
        else:
            box((560,165,960,330),'Order service',['POST /orders; GET /orders/{id}','Persistent client idempotency key','One process / one demo Pod'])
            box((560,470,960,635),'Inventory service',['GET /books; POST /reservations','Atomic reservation + deduction','One process / one demo Pod'])
            box((1070,165,1450,330),'Order database',['SQLite on own PVC','Orders only'],'#d9f2e5')
            box((1070,470,1450,635),'Inventory database',['SQLite on separate PVC','Books + reservations'],'#d9f2e5')
            arrow((400,235),(560,235),'REST')
            draw.line([(225,300),(225,540)],fill='#386393',width=4)
            arrow((225,540),(560,540),'GET /books')
            arrow((760,330),(760,470),'HTTP + service token')
            arrow((960,235),(1070,235))
            arrow((960,540),(1070,540))
            box((50,715,1450,865),'Deployment boundary',['Local processes OR Docker Compose OR Minikube namespace: bookflow',
                 'ClusterIP Services + Secret + health probes + dedicated PVCs; replicas = 1; update strategy = Recreate',
                 'No cross-service database access. Production scaling requires external databases and a hardened runtime.'],'#edf0f7')
        name = 'scalable_design' if production else 'architecture'
        canvas.save(folder / f'{STUDENT}_{name}.png')

    datafile = folder / 'api-results.json'
    if datafile.exists():
        records = json.loads(datafile.read_text(encoding='utf-8'))
        selected = ['Create order','Safe retry','Stock deducted once','Concurrency summary','Dependency outage','Recovery with same key']
        panel = Image.new('RGB', (1500, 1490), '#101c2e')
        draw = ImageDraw.Draw(panel)
        draw.text((40,25),'BookFlow | actual local API results',font=font(36),fill='#f0f6ff')
        draw.text((40,80),'Rendered from api-results.json; not Postman / Kubernetes screenshots.',font=font(23),fill='#a9c6e9')
        for i,label in enumerate(selected):
            record = next(r for r in records if r['label'] == label)
            x,y = 35+(i%2)*740, 145+(i//2)*440
            draw.rounded_rectangle((x,y,x+710,y+415),radius=12,fill='#20314a')
            draw.text((x+18,y+16),label,font=font(26),fill='#8dded0')
            body = {k:v for k,v in record.items() if k not in ['label','request']}
            lines = []
            for line in json.dumps(body, indent=2).splitlines():
                lines.extend(textwrap.wrap(line, width=53, replace_whitespace=False) or [''])
            # No data is omitted in the source JSON; panel shows the key fields compactly.
            if len(lines) > 17:
                body = {k:v for k,v in body.items() if k != 'url'}
                lines = json.dumps(body,indent=2).splitlines()
            for j,line in enumerate(lines):
                draw.text((x+18,y+62+j*20),line,font=font(17),fill='#eef4ff')
        panel.save(folder / f'{STUDENT}_api_evidence.png')


def service_readmes():
    for service, port in [('inventory-service',8001),('order-service',8002)]:
        inventory = service.startswith('inventory')
        description = ('Owns books and reservations. POST /reservations requires X-Service-Token. '
                       'Concurrent stock checks and deductions are atomic; replay by order_id does not deduct twice.' if inventory else
                       'Owns orders and client idempotency keys. POST /orders requires Idempotency-Key. '
                       'Calls Inventory with a two-second timeout; retry the same key after HTTP 503.')
        routes = ('GET /books; GET /books/{sku}; POST /reservations' if inventory else
                  'POST /orders; GET /orders/{id}; GET /orders (up to 100 rows)')
        write(f'{service}/README.md', f'''# BookFlow {service}

Student: {STUDENT} — {NAME} (individual), SEZG583.

{description}

## Run independently

Python 3.13; standard library only. No pip installation is needed.

```powershell
$env:SERVICE_TOKEN = 'replace-with-the-same-random-token-on-both-services'
$env:PORT = '{port}'
$env:HOST = '127.0.0.1'
$env:DB_PATH = 'data/service.db'
''' + ("$env:INVENTORY_URL = 'http://127.0.0.1:8001'\n" if not inventory else '') + f'''python app.py
```

Routes: {routes}. Both services expose GET /health/live and /health/ready.

```powershell
python -m unittest discover -s tests -v
docker build -t bookflow-{'inventory' if inventory else 'orders'}:1.0 .
```

The Order unit suite mocks the HTTP dependency; it runs without a live Inventory
service. The Inventory suite uses isolated temporary SQLite databases.
Containers listen on port 8000, run as UID 10001, and store data at /data/service.db.
Pass SERVICE_TOKEN with -e and mount a dedicated /data volume when running the image.
Order additionally needs an INVENTORY_URL reachable on the container network.
No files or modules from a sibling repository are needed to build or test this service.

CI runs unit tests and a Docker build independently. No remote CI success or CD is
claimed. A single-process SQLite deployment is a teaching configuration; production
scaling requires an external database, hardened HTTP runtime, authorization and TLS.
Do not expose these assessment APIs to the public internet.
''')


def docx_report():
    source = ROOT / f'{STUDENT}_BookFlow_Report.md'
    if not source.exists():
        write(source.name, REPORT)
    document = Document()
    document.core_properties.author = NAME
    document.core_properties.title = 'BookFlow — Microservices Development and Deployment'
    section = document.sections[0]
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header.paragraphs[0].text = 'SEZG583  |  BookFlow  |  2025m203021'
    section.footer.paragraphs[0].text = f'{NAME}  |  Individual submission'
    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    table_lines = []
    code_lines = []
    in_code = False

    def flush_table():
        if not table_lines:
            return
        rows = [[cell.strip() for cell in line.strip('|').split('|')] for line in table_lines]
        rows = [row for row in rows if not all(set(cell) <= set('-: ') for cell in row)]
        table = document.add_table(rows=0, cols=len(rows[0]))
        table.style = 'Light Shading Accent 1'
        for row in rows:
            cells = table.add_row().cells
            for i, cell in enumerate(row):
                cells[i].text = cell
        document.add_paragraph()
        table_lines.clear()

    paragraph_lines = []
    def flush_paragraph():
        if paragraph_lines:
            document.add_paragraph(' '.join(paragraph_lines))
            paragraph_lines.clear()

    for line in source.read_text(encoding='utf-8').splitlines():
        if line.startswith('```'):
            flush_paragraph()
            flush_table()
            if in_code:
                p = document.add_paragraph('\n'.join(code_lines))
                for run in p.runs:
                    run.font.name = 'Consolas'
                    run.font.size = Pt(8)
                code_lines.clear()
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith('|'):
            flush_paragraph()
            table_lines.append(line)
            continue
        flush_table()
        if not line.strip():
            flush_paragraph()
        elif line.startswith('#'):
            flush_paragraph()
            level = len(line) - len(line.lstrip('#'))
            document.add_heading(line.lstrip('# '), level=min(level - 1, 3))
        elif line.startswith('!['):
            flush_paragraph()
            label, path = line[2:].split('](',1)
            image = ROOT / path.rstrip(')')
            if image.exists():
                document.add_picture(str(image), width=Inches(6.2))
                document.add_paragraph(label, style='Caption')
        elif line.startswith('- '):
            flush_paragraph()
            document.add_paragraph(line[2:],style='List Bullet')
        else:
            paragraph_lines.append(line)
    flush_paragraph()
    flush_table()
    document.save(ROOT / f'{STUDENT}_BookFlow_Report.docx')


def package():
    destination = ROOT / f'{STUDENT}_BookFlow.zip'
    with zipfile.ZipFile(destination, 'w', zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(ROOT.rglob('*')):
            if not path.is_file() or path == destination:
                continue
            relative = path.relative_to(ROOT)
            if any(part in {'.git','__pycache__','data','.venv'} for part in relative.parts):
                continue
            if path.suffix in {'.zip','.db','.pyc'} or path.name.endswith(('.db-wal','.db-shm')):
                continue
            archive.write(path, str(Path(f'{STUDENT}_BookFlow') / relative))
    print(f'Created {destination.name} ({destination.stat().st_size:,} bytes)')


def main():
    manifests()
    postman()
    write('README.md', README)
    write('tools/requirements-docs.txt', 'python-docx>=1.1,<2\nPillow>=11,<13')
    service_readmes()
    diagrams()
    docx_report()
    package()


if __name__ == '__main__':
    main()
